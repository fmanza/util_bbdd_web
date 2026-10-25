"""Retoque del Excel de la documentación con una plantilla (ADR-011, GAPS G40).

util_bbdd 4.0.0 no admite plantilla Excel: ``ExcelGenerator`` crea un libro nuevo con colores RGB fijos (la paleta de
Office) y fuentes sin nombre. Después de generarlo, la web pone el tema de la plantilla, cambia esos colores por los
del tema (``accent1``/``accent5`` con su tono, que con el tema de Office se ven igual), hace que las fuentes sigan la
fuente del tema y copia el encabezado, el pie y la configuración de página de la primera hoja de la plantilla. Todas
las tablas toman el estilo predefinido elegido. Sin portada, sin imágenes en el encabezado ni estilos de tabla propios.

La tabla de colores depende de util_bbdd 4.0.0: la protege un test que documenta una base SQLite real.
"""

from __future__ import annotations

import copy
import re
from pathlib import Path
from typing import Any, Final

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import Color

#: Colores fijos de ``ExcelGenerator`` (util_bbdd 4.0.0) → (índice de color del tema, tono).
#: Índices del tema en Excel: 4 = ``accent1``, 8 = ``accent5``.
FIXED_COLORS: Final[dict[str, tuple[int, float]]] = {
    "4472C4": (4, 0.0),  # accent1 de Office 2013+
    "4F81BD": (4, 0.0),  # accent1 de Office 2007
    "366092": (4, -0.25),  # accent1 de Office 2007, 25 % más oscuro
    "5B9BD5": (8, 0.0),  # accent5 de Office 2013+
    "2F75B5": (8, -0.25),  # accent5, 25 % más oscuro
    "D9E1F2": (4, 0.8),  # accent1, 80 % más claro (bandas de las filas)
}

#: Estilos de tabla predefinidos de Excel (los únicos que se pueden elegir: G40).
TABLE_STYLES: Final[tuple[str, ...]] = (
    *(f"TableStyleLight{i}" for i in range(1, 22)),
    *(f"TableStyleMedium{i}" for i in range(1, 29)),
    *(f"TableStyleDark{i}" for i in range(1, 12)),
)
DEFAULT_TABLE_STYLE: Final = "TableStyleMedium2"

_THEME_RELS = re.compile(rb"\br:(?:embed|link|id)\s*=")
_MINOR_FONT = re.compile(rb"<a:minorFont>\s*<a:latin\s+typeface=\"([^\"]*)\"")
_IMAGE_CODE = re.compile(r"&G")
_HEADER_PARTS: Final = ("oddHeader", "oddFooter", "evenHeader", "evenFooter", "firstHeader", "firstFooter")


class ExcelTemplateError(ValueError):
    """La plantilla Excel no se puede aplicar (sin hojas o con un tema que referencia otros ficheros)."""


def theme_is_safe(theme: bytes | None) -> bool:
    """Si el tema se puede copiar tal cual: openpyxl no copia sus ``.rels`` (imágenes), y Excel lo daría por dañado."""
    return theme is None or _THEME_RELS.search(theme) is None


def _theme_color(color: Any) -> Color | None:
    """Color del tema equivalente a un color fijo conocido (``None`` si no es uno de ellos)."""
    rgb = getattr(color, "rgb", None)
    if getattr(color, "type", None) != "rgb" or not isinstance(rgb, str):
        return None
    mapped = FIXED_COLORS.get(rgb[-6:].upper())
    return Color(theme=mapped[0], tint=mapped[1]) if mapped else None


def _recolor(cell: Any, minor_font: str | None) -> None:
    """Rellenos con colores fijos → colores del tema; fuentes sin nombre → la fuente del tema."""
    if isinstance(cell, MergedCell):
        return
    fill = cell.fill
    if fill.fill_type == "solid" and (color := _theme_color(fill.fgColor)) is not None:
        cell.fill = PatternFill(fill_type="solid", fgColor=color, bgColor=copy.copy(color))
    font = cell.font
    if minor_font and font.scheme is None and font.name in (None, "Calibri"):
        new = copy.copy(font)
        new.name, new.scheme = minor_font, "minor"
        cell.font = new


def _header_footer(source: Any, target: Any) -> None:
    """Encabezado, pie y página de la hoja de la plantilla, sin imágenes (``&G``)."""
    for part in _HEADER_PARTS:
        header = copy.deepcopy(getattr(source, part))
        for side in (header.left, header.center, header.right):
            if side.text:
                side.text = _IMAGE_CODE.sub("", side.text) or None
        setattr(target, part, header)
    target.HeaderFooter.differentFirst = source.HeaderFooter.differentFirst
    target.HeaderFooter.differentOddEven = source.HeaderFooter.differentOddEven
    for attr in ("orientation", "paperSize", "fitToWidth", "fitToHeight", "scale"):
        setattr(target.page_setup, attr, getattr(source.page_setup, attr))
    target.sheet_properties.pageSetUpPr = copy.copy(source.sheet_properties.pageSetUpPr)
    target.page_margins = copy.copy(source.page_margins)
    target.print_options = copy.copy(source.print_options)


def _read_template(template: Path) -> tuple[bytes | None, Any]:
    """Tema (si es seguro copiarlo) y primera hoja de la plantilla.

    Raises:
        ExcelTemplateError: Sin hojas o con un tema que referencia otros ficheros.
    """
    source = load_workbook(template)
    theme: bytes | None = source.loaded_theme
    if not theme_is_safe(theme):
        raise ExcelTemplateError("El tema de la plantilla referencia imágenes u otros ficheros.")
    if not source.worksheets:
        raise ExcelTemplateError("La plantilla no tiene hojas.")
    return theme, source.worksheets[0]


def _retouch_sheet(ws: Any, sheet: Any, minor_font: str | None, style: str) -> None:
    for row in ws.iter_rows():
        for cell in row:
            _recolor(cell, minor_font)
    for table in ws.tables.values():
        if table.tableStyleInfo is not None:
            table.tableStyleInfo.name = style
    _header_footer(sheet, ws)


def apply_excel_template(book: Path, template: Path, table_style: str | None) -> None:
    """Aplica ``template`` a ``book`` en el sitio; si algo falla, ``book`` queda como estaba.

    Raises:
        ExcelTemplateError: La plantilla no tiene hojas o su tema referencia otros ficheros.
        Las de openpyxl si alguno de los dos libros no se puede leer.
    """
    theme, sheet = _read_template(template)
    match = _MINOR_FONT.search(theme) if theme else None
    minor_font = match.group(1).decode("utf-8", "replace") if match and match.group(1) else None
    style = table_style if table_style in TABLE_STYLES else DEFAULT_TABLE_STYLE

    workbook = load_workbook(book)
    if theme is not None:
        workbook.loaded_theme = theme
    for ws in workbook.worksheets:
        _retouch_sheet(ws, sheet, minor_font, style)
    tmp = book.with_name(f".tmp-{book.name}")
    try:
        workbook.save(tmp)
        tmp.replace(book)
    finally:
        tmp.unlink(missing_ok=True)
