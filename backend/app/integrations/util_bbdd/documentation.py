"""Documentación y grafos de relaciones con util_bbdd (Fase 5); se ejecutan dentro del subproceso del job.

- ``DatabaseDocumentation(manager, directorio_salida, fichero_salida, imprimir, formatos, batch_window_seconds=0)``:
  siempre con ``fichero_salida`` explícito y sin ventana de agrupación (GAPS G2; desde util_bbdd 5.0.0 la ventana por
  defecto ya es 0, pero se sigue pasando). Su salida
  (``<dir>/<nombre>/{WORD,EXCEL,HTML}``) se recoloca en ``outputs/{word,excel,html}/`` del run.
- ``GraphvizRelationshipGraphExt``: grafo paginado de una tabla o de todas las que tienen relaciones.

El schema se pasa como argumento (``schema=``, util_bbdd 4.1.0; GAPS G22): en PostgreSQL, SQL Server y HSQLDB es el
schema; en MySQL/MariaDB, la base de datos; en Oracle, el owner (también otro que el usuario, si tiene permiso de
lectura; GAPS G21). En SQLite no aplica (solo existe ``main``).

Plantillas (ADR-011): van en ``imprimir["plantilla_word"]`` y ``imprimir["plantilla_excel"]`` (util_bbdd 4.1.0, G40 y
G41), que las valida al construir ``DatabaseDocumentation``. Si una falla, se genera de nuevo sin esa plantilla y la
ejecución termina con un aviso en lugar de un error.
"""

from __future__ import annotations

import logging
import shutil
import subprocess
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Final, Literal

from docx.opc.exceptions import PackageNotFoundError

from app.integrations.util_bbdd.progress import RunProgress

logger = logging.getLogger(__name__)

Direction = Literal["both", "outgoing", "incoming"]
#: Columnas de los diagramas (documentación y grafo de relaciones). En la web ``compact`` es «sin columnas»; el
#: ``compact`` de util_bbdd (primero las de relación, luego el resto) es ``relations_first``.
GraphColumns = Literal["relations_first", "relations", "all", "compact"]
_GRAPH_COLUMN_MODE: Final[dict[str, str]] = {"relations_first": "compact", "relations": "relations", "all": "all"}

#: Procesos ``dot`` simultáneos por job (util_bbdd usa por defecto todos los núcleos hasta 8).
GRAPH_WORKERS: Final = 2
_FORMAT_DIRS: Final[dict[str, str]] = {"word": "WORD", "excel": "EXCEL", "html": "HTML"}
#: Estilos que ``WordGenerator`` usa con nombre (``add_heading`` 0-3 y el DDL): si la plantilla no los tiene, python-docx
#: lanza ``KeyError`` en mitad del documento (GAPS G41).
WORD_REQUIRED_STYLES: Final = ("Title", "Heading 1", "Heading 2", "Heading 3", "No Spacing")


class GraphvizMissingError(RuntimeError):
    """Graphviz (``dot``) no está instalado."""


def graphviz_failed(exc: BaseException) -> bool:
    """Si ``exc`` viene de un ``dot`` que terminó con error (``graphviz.CalledProcessError``).

    util_bbdd solo captura ``OSError`` y ``RuntimeError`` por diagrama: el fallo de uno llega envuelto en el error del
    formato entero (``DocumentationRenderError``), con el de Graphviz en la cadena de causas (GAPS G39).
    """
    seen: set[int] = set()
    current: BaseException | None = exc
    while current is not None and id(current) not in seen:
        if isinstance(current, subprocess.CalledProcessError):
            return True
        seen.add(id(current))
        current = current.__cause__ or current.__context__
    return False


def target_schema(db_type: str, schema: str | None) -> str | None:
    """Schema que se pasa a util_bbdd: ``None`` si no se eligió o en SQLite (solo ``main``)."""
    if db_type == "sqlite" or not schema or not schema.strip():
        return None
    return schema.strip()


@dataclass(frozen=True)
class DocumentationOptions:
    """Opciones de la operación ``documentation`` (las del formulario, sin rutas del servidor)."""

    formats: tuple[str, ...]
    tables: bool = True
    table_columns: bool = True
    table_keys: bool = True
    table_indexes: bool = True
    table_ddl: bool = False
    views: bool = True
    view_ddl: bool = False
    routines: bool = False
    database_info: bool = True
    storage: bool = False
    engine_conditions: bool = False
    charset: bool = False
    graphs: bool = True
    graph_depth: int = 1
    graph_direction: Direction = "both"
    graph_columns: GraphColumns = "relations_first"
    graph_max_columns: int = 8
    graph_data_types: bool = False
    graph_orthogonal: bool = True
    tables_per_document: int = 50
    html_items_per_page: int = 15


@dataclass(frozen=True)
class DocumentTemplates:
    """Plantillas de una documentación: copias en la carpeta del run (nunca rutas que elija el usuario)."""

    word: Path | None = None
    word_table_style: str | None = None
    excel: Path | None = None
    excel_table_style: str | None = None


@dataclass
class _Outcome:
    warnings: list[str] = field(default_factory=list)
    applied: dict[str, bool] = field(default_factory=dict)


def word_template_failed(exc: BaseException) -> bool:
    """Si ``exc`` viene de la plantilla Word: estilo ausente (G41), fichero que no es un Word o plantilla no encontrada.

    Los fallos de Graphviz no cuentan: generar sin plantilla no los arregla.
    """
    from util_bbdd.exceptions import DocumentationConfigError

    if graphviz_failed(exc):
        return False
    seen: set[int] = set()
    current: BaseException | None = exc
    while current is not None and id(current) not in seen:
        message = str(current)
        if (
            (isinstance(current, KeyError) and "style" in message)
            or isinstance(current, zipfile.BadZipFile)
            or isinstance(current, PackageNotFoundError)
            or (isinstance(current, ValueError) and "not a Word file" in message)
            # util_bbdd 4.1.0 valida la plantilla al construir: «A la plantilla Word … le faltan estilos…»
            or (isinstance(current, DocumentationConfigError) and "plantilla word" in message.lower())
        ):
            return True
        seen.add(id(current))
        current = current.__cause__ or current.__context__
    return False


def excel_template_failed(exc: BaseException) -> bool:
    """Si ``exc`` viene de la plantilla Excel (util_bbdd 4.1.0: no existe, no es un libro, sin hojas o con un tema que
    referencia otros ficheros; ``DocumentationConfigError`` «… plantilla Excel …»)."""
    from util_bbdd.exceptions import DocumentationConfigError

    seen: set[int] = set()
    current: BaseException | None = exc
    while current is not None and id(current) not in seen:
        if isinstance(current, DocumentationConfigError) and "plantilla excel" in str(current).lower():
            return True
        seen.add(id(current))
        current = current.__cause__ or current.__context__
    return False


def build_imprimir(
    options: DocumentationOptions,
    *,
    word_template: Path | None = None,
    table_style: str | None = None,
    excel_template: Path | None = None,
    excel_table_style: str | None = None,
) -> dict[str, Any]:
    """Configuración ``imprimir`` de ``DatabaseDocumentation`` (``PrintConfigManager``).

    Las plantillas son copias en la carpeta del run; ``table_style`` es el estilo de tabla de Word y
    ``excel_table_style``, el del Excel (``table_style_excel``, util_bbdd 4.1.0).
    """
    extra: dict[str, Any] = {}
    if word_template is not None:
        extra["plantilla_word"] = str(word_template)
        if table_style:
            extra["table_style"] = table_style
    if excel_template is not None:
        extra["plantilla_excel"] = str(excel_template)
        if excel_table_style:
            extra["table_style_excel"] = excel_table_style
    return {
        **extra,
        "tablas": {
            "imprimir_tablas": options.tables,
            "columnas": options.tables and options.table_columns,
            "pk": options.tables and options.table_keys,
            "fk": options.tables and options.table_keys,
            "unicas": options.tables and options.table_keys,
            "indices": options.tables and options.table_indexes,
            "ddl": options.tables and options.table_ddl,
            "charset_collation": options.tables and options.charset,
        },
        "vistas": {
            "imprimir_vistas": options.views,
            "columnas": options.views,
            "ddl": options.views and options.view_ddl,
        },
        "rutinas": {"imprimir_procedures": options.routines, "imprimir_function": options.routines},
        "almacenamiento": {"imprimir_almacenamiento": options.storage},
        "condiciones_motor": {"imprimir_condiciones_motor": options.engine_conditions},
        "info_database": {
            "imprimir_info_database": options.database_info,
            "version": options.database_info,
            "charset_collation": options.database_info and options.charset,
        },
        "columnas_charset": {"charset_collation": options.charset},
        "grafo": {
            "enabled": options.graphs,
            "max_depth": options.graph_depth,
            "direction": options.graph_direction,
            "workers": GRAPH_WORKERS,
            "show_columns": options.graph_columns != "compact",
            "orthogonal_edges": options.graph_orthogonal,  # util_bbdd 4.1.0 (G39)
            "columns": {
                "mode": _GRAPH_COLUMN_MODE.get(options.graph_columns, "relations"),
                "max_items": options.graph_max_columns,
                "show_data_types": options.graph_data_types,
            },
        },
        "tablas_por_documento": options.tables_per_document,
        "html_items_por_pagina": options.html_items_per_page,
    }


def _move(source: Path, target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        shutil.rmtree(target) if target.is_dir() else target.unlink()
    shutil.move(str(source), str(target))


def generate_documentation(
    db_type: str,
    conn_params: dict[str, Any],
    *,
    schema: str | None,
    options: DocumentationOptions,
    outputs_dir: Path,
    work_dir: Path,
    name: str,
    templates: DocumentTemplates | None = None,
    progress: RunProgress | None = None,
) -> dict[str, Any]:
    """Genera la documentación y la deja en ``outputs_dir/{word,excel,html/documentacion}``.

    Con ``progress`` (GAPS G1, util_bbdd 4.1.0) el run muestra el avance por formato y tabla y se puede cancelar entre
    tablas (``OperationCancelledError``).

    Returns:
        ``{"tables", "views", "routines", "formats": {formato: ruta relativa a outputs_dir}, "warnings": [...],
        "templates_applied": {"word"|"excel": bool}}`` (solo los formatos con plantilla).

    Raises:
        GraphvizMissingError y las excepciones de util_bbdd.
    """
    if options.graphs and shutil.which("dot") is None:
        raise GraphvizMissingError("Graphviz (dot) no está instalado en el servidor.")
    chosen = templates or DocumentTemplates()
    outcome = _Outcome()
    try:
        summary = _document_with_template(
            conn_params,
            db_type,
            options,
            outputs_dir,
            work_dir,
            name,
            chosen,
            outcome,
            target_schema(db_type, schema),
            progress,
        )
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)
    return {**summary, "warnings": outcome.warnings, "templates_applied": outcome.applied}


def _document_with_template(
    params: dict[str, Any],
    db_type: str,
    options: DocumentationOptions,
    outputs_dir: Path,
    work_dir: Path,
    name: str,
    templates: DocumentTemplates,
    outcome: _Outcome,
    schema: str | None,
    progress: RunProgress | None = None,
) -> dict[str, Any]:
    """Documenta con las plantillas; si falla por una (G40, G41), se repite sin ella (como mucho, sin las dos)."""
    word = templates.word if "word" in options.formats else None
    excel = templates.excel if "excel" in options.formats else None
    while True:
        imprimir = build_imprimir(
            options,
            word_template=word,
            table_style=templates.word_table_style,
            excel_template=excel,
            excel_table_style=templates.excel_table_style,
        )
        try:
            summary = _document(params, db_type, options, outputs_dir, work_dir, name, imprimir, schema, progress)
        except Exception as exc:  # frontera de util_bbdd: solo se reintenta si el fallo es de una plantilla
            if word is not None and word_template_failed(exc):
                word = None
                _fallback(outcome, "word", exc)
            elif excel is not None and excel_template_failed(exc):
                excel = None
                _fallback(outcome, "excel", exc)
            else:
                raise
            shutil.rmtree(work_dir, ignore_errors=True)
            continue
        if word is not None:
            outcome.applied["word"] = True
        if excel is not None:
            outcome.applied["excel"] = True
        return summary


def _fallback(outcome: _Outcome, kind: Literal["word", "excel"], exc: BaseException) -> None:
    label = "Word" if kind == "word" else "Excel"
    logger.warning(f"Plantilla {label} no aplicada", extra={"template_fallback": kind, "exc_type": type(exc).__name__})
    outcome.applied[kind] = False
    target = "el documento Word" if kind == "word" else "el libro Excel"
    outcome.warnings.append(f"La plantilla {label} no se pudo aplicar: {target} usa el formato por defecto.")


def _document(
    params: dict[str, Any],
    db_type: str,
    options: DocumentationOptions,
    outputs_dir: Path,
    work_dir: Path,
    name: str,
    imprimir: dict[str, Any],
    schema: str | None,
    progress: RunProgress | None = None,
) -> dict[str, Any]:
    from util_bbdd import SyncDatabaseManager
    from util_bbdd.doc import DatabaseDocumentation

    with SyncDatabaseManager(db_type, params, use_pool=False) as manager:
        result = DatabaseDocumentation(
            manager,
            directorio_salida=work_dir,
            fichero_salida=Path(name),
            imprimir=imprimir,
            formatos=list(options.formats),
            batch_window_seconds=0,
            schema=schema,
            on_progress=progress,
            cancel_token=progress.token if progress is not None else None,
        ).generate()
    produced: dict[str, str] = {}
    base = work_dir / name
    for fmt in options.formats:
        folder = base / _FORMAT_DIRS[fmt]
        if not folder.exists():
            continue
        if fmt == "html":
            target = outputs_dir / "html" / "documentacion"
            _move(folder, target)
            produced[fmt] = (target / "index.html").relative_to(outputs_dir).as_posix()
        else:
            files = sorted(p for p in folder.iterdir() if p.is_file())
            for file in files:
                _move(file, outputs_dir / fmt / file.name)
            if files:
                produced[fmt] = (outputs_dir / fmt / files[0].name).relative_to(outputs_dir).as_posix()
    return {"tables": result.tables, "views": result.views, "routines": result.routines, "formats": produced}


@dataclass(frozen=True)
class GraphOptions:
    """Opciones de la operación ``relationship_graph``."""

    table: str | None = None
    max_depth: int = 1
    direction: Direction = "both"
    column_mode: GraphColumns = "relations"
    max_columns: int = 8
    show_data_types: bool = False
    orthogonal_edges: bool = True
    max_subnodes: int = 10
    format: Literal["svg", "png"] = "svg"


def generate_graphs(
    db_type: str,
    conn_params: dict[str, Any],
    *,
    schema: str | None,
    options: GraphOptions,
    output_dir: Path,
    progress: RunProgress | None = None,
) -> dict[str, Any]:
    """Grafo de ``options.table`` (paginado si hay muchas relaciones) o de todas las tablas con relaciones.

    Con ``progress`` (GAPS G1, util_bbdd 4.1.0), el grafo de todas las tablas avisa tabla a tabla y se puede cancelar
    entre tablas (``OperationCancelledError``); el de una tabla es una sola unidad.

    Returns:
        ``{"files": [nombres], "tables": n, "errors": n}``.

    Raises:
        GraphvizMissingError, LookupError (tabla inexistente) y las de util_bbdd.
    """
    from util_bbdd import SyncDatabaseManager
    from util_bbdd.analysis.graphviz_ext import GraphvizRelationshipGraphExt

    if shutil.which("dot") is None:
        raise GraphvizMissingError("Graphviz (dot) no está instalado en el servidor.")
    graph_schema = target_schema(db_type, schema)
    output_dir.mkdir(parents=True, exist_ok=True)
    with SyncDatabaseManager(db_type, conn_params, use_pool=False) as manager:
        graph = GraphvizRelationshipGraphExt(
            manager,
            show_columns=options.column_mode != "compact",
            column_mode=_GRAPH_COLUMN_MODE.get(options.column_mode, "relations"),
            max_columns=options.max_columns,
            show_data_types=options.show_data_types,
            orthogonal_edges=options.orthogonal_edges,
            max_depth=options.max_depth,
            direction=options.direction,
            render_workers=GRAPH_WORKERS,
            schema=graph_schema,
        )
        if options.table:
            tables = {t.name for t in manager.get_schema_metadata(graph_schema, kinds={_tables_kind()}).tables()}
            if options.table not in tables:
                raise LookupError(f"La tabla «{options.table}» no existe en el schema.")
            files = graph.generate_graphviz_graph_paginated_ext(
                table_name=options.table,
                output_dir=output_dir,
                format=options.format,
                max_subnodes=options.max_subnodes,
            )
            return {"files": [p.name for p in files], "tables": 1, "errors": 0}
        results = graph.generate_all_tables_graphs_ext(
            output_dir=output_dir,
            format=options.format,
            max_subnodes=options.max_subnodes,
            on_progress=progress,
            cancel_token=progress.token if progress is not None else None,
        )
    names = [Path(f).name for r in results.values() for f in r.get("files", [])]
    errors = sum(1 for r in results.values() if r.get("status") == "error")
    return {"files": names, "tables": len(results), "errors": errors}


def _tables_kind() -> Any:
    from util_bbdd.metadata import Kind

    return Kind.TABLES
