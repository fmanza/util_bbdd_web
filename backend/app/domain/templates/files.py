"""Comprobación de las plantillas de documento (ADR-011) y plantillas base, sin util_bbdd.

Una plantilla de un usuario acaba en los documentos que descargan otros: además de comprobar que es un Word o un Excel
de verdad (ZIP con su parte principal, tamaño descomprimido acotado), se rechaza lo que Word o Excel podrían ejecutar
o descargar al abrirlo: macros, objetos incrustados (OLE, ActiveX), relaciones externas que no sean hipervínculos
(plantilla adjunta remota, imágenes enlazadas…), campos DDE/INCLUDE/LINK y cualquier DTD (``<!DOCTYPE``: OOXML no la
usa; sin ella no hay entidades). Los ``.dotx``/``.xltx`` se guardan como ``.docx``/``.xlsx`` (python-docx no abre
plantillas). Word: deben estar los estilos que usa util_bbdd (GAPS G41).
"""

from __future__ import annotations

import html
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePath
from typing import Final, Literal

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from openpyxl import Workbook, load_workbook

from app.domain.templates.errors import TemplateContentError, TemplateTypeError
from app.integrations.util_bbdd.documentation import WORD_REQUIRED_STYLES
from app.integrations.util_bbdd.excel_template import DEFAULT_TABLE_STYLE, TABLE_STYLES, theme_is_safe

Kind = Literal["word", "excel"]
KINDS: Final[tuple[Kind, ...]] = ("word", "excel")
#: Extensión → (tipo, es una plantilla de Office que hay que convertir).
EXTENSIONS: Final[dict[str, tuple[Kind, bool]]] = {
    ".docx": ("word", False),
    ".dotx": ("word", True),
    ".xlsx": ("excel", False),
    ".xltx": ("excel", True),
}
STORED_SUFFIX: Final[dict[str, str]] = {"word": "docx", "excel": "xlsx"}
_MACRO_EXTENSIONS: Final = frozenset({".docm", ".dotm", ".xlsm", ".xltm"})
_MAIN_PART: Final[dict[Kind, str]] = {"word": "word/document.xml", "excel": "xl/workbook.xml"}
_TEMPLATE_TYPES: Final[dict[Kind, tuple[bytes, bytes]]] = {
    "word": (b"wordprocessingml.template.main+xml", b"wordprocessingml.document.main+xml"),
    "excel": (b"spreadsheetml.template.main+xml", b"spreadsheetml.sheet.main+xml"),
}
_EMBEDDED: Final = ("/embeddings/", "/activex/", "vbaproject.bin", "/externallinks/")
_RELATIONSHIP = re.compile(r"<(?:\w+:)?Relationship\b([^>]*)>")
_ATTRIBUTE = re.compile(r"([\w:]+)\s*=\s*(?:\"([^\"]*)\"|'([^']*)')")
#: Piezas de los campos complejos de Word en orden: ``fldChar`` (inicio, separador, fin) y texto del código
#: (``instrText``, que Word puede partir en varios trozos).
_FIELD_TOKEN = re.compile(
    r"<(?:\w+:)?fldChar\b[^>]*?\bfldCharType\s*=\s*[\"'](begin|separate|end)[\"']"
    r"|<(?:\w+:)?instrText\b[^>]*(?<!/)>([^<]*)"
)
_INSTR_ATTR = re.compile(r"\b(?:\w+:)?instr\s*=\s*(?:\"([^\"]*)\"|'([^']*)')")
#: Tipos de campo que leen o ejecutan algo de fuera al abrir o actualizar el documento (la primera palabra del código).
_DANGEROUS_FIELDS: Final = frozenset({"DDE", "DDEAUTO", "INCLUDE", "INCLUDETEXT", "INCLUDEPICTURE", "LINK", "IMPORT"})
_CONTROL = re.compile(r"[\x00-\x1f\x7f]")


@dataclass(frozen=True)
class TemplateInfo:
    """Lo que se guarda de una plantilla comprobada: estilos de tabla elegibles y el inicial."""

    table_styles: list[str]
    table_style: str | None


def display_name(filename: str | None) -> str:
    """Nombre que se muestra: solo la última parte, sin caracteres de control y acotado (nunca es una ruta)."""
    name = PurePath((filename or "").replace("\\", "/")).name
    return _CONTROL.sub("", name).strip()[:255] or "plantilla"


def kind_for(filename: str) -> tuple[Kind, bool]:
    """Tipo de plantilla por la extensión y si hay que convertirla (``.dotx``/``.xltx``).

    Raises:
        TemplateTypeError: Extensión no admitida (las de macros, con su propio mensaje).
    """
    suffix = PurePath(filename).suffix.lower()
    if suffix in _MACRO_EXTENSIONS:
        raise TemplateTypeError("No se admiten plantillas con macros.")
    found = EXTENSIONS.get(suffix)
    if found is None:
        raise TemplateTypeError()
    return found


def untemplate(path: Path, kind: Kind) -> None:
    """Convierte en el sitio un ``.dotx``/``.xltx`` en documento (solo cambia el tipo de la parte principal).

    Raises:
        TemplateContentError: No es un ZIP.
    """
    if not zipfile.is_zipfile(path):
        raise TemplateContentError("La plantilla no es un documento de Office (no es un ZIP).")
    source, target = _TEMPLATE_TYPES[kind]
    buffer = io.BytesIO()
    with zipfile.ZipFile(path) as archive, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as out:
        for info in archive.infolist():
            data = archive.read(info.filename)
            if info.filename == "[Content_Types].xml":
                data = data.replace(source, target)
            out.writestr(info, data)
    path.write_bytes(buffer.getvalue())


def _attributes(raw: str) -> dict[str, str]:
    return {m.group(1).split(":")[-1]: html.unescape(m.group(2) or m.group(3) or "") for m in _ATTRIBUTE.finditer(raw)}


@dataclass
class _Field:
    code: str = ""
    open: bool = True  # recogiendo el código (hasta el separador o el fin)
    computed: bool = False  # el tipo lo da un campo anidado: no se puede saber cuál es


class _FieldCodes:
    """Reconstruye el código de cada campo complejo a partir de sus piezas, en orden y con anidamiento."""

    def __init__(self) -> None:
        self.codes: list[str | None] = []
        self._stack: list[_Field] = []

    def _close(self, field: _Field) -> None:
        if field.open:
            field.open = False
            self.codes.append(None if field.computed else field.code)

    def begin(self) -> None:
        if self._stack and self._stack[-1].open and not self._stack[-1].code.strip():
            self._stack[-1].computed = True  # un campo anidado antes del tipo: el tipo se calcula
        self._stack.append(_Field())

    def separate(self) -> None:
        if self._stack:
            self._close(self._stack[-1])

    def end(self) -> None:
        if self._stack:
            self._close(self._stack.pop())

    def text(self, instr: str) -> None:
        if not self._stack:  # ``instrText`` sin inicio: cuenta como un campo
            self._stack.append(_Field())
        if self._stack[-1].open:
            self._stack[-1].code += html.unescape(instr)

    def finish(self) -> list[str | None]:
        for field in self._stack:
            self._close(field)
        return self.codes


def _field_codes(text: str) -> list[str | None]:
    """Código de cada campo de una parte XML (``None`` si su tipo lo calcula un campo anidado).

    Los ``instrText`` sin ``fldChar`` de inicio cuentan como un campo; los ``fldSimple`` dan su ``instr``.
    """
    fields = _FieldCodes()
    for match in _FIELD_TOKEN.finditer(text):
        kind, instr = match.group(1), match.group(2)
        if kind == "begin":
            fields.begin()
        elif kind == "separate":
            fields.separate()
        elif kind == "end":
            fields.end()
        elif instr is not None:
            fields.text(instr)
    simple = [html.unescape(m.group(1) or m.group(2) or "") for m in _INSTR_ATTR.finditer(text)]
    return [*fields.finish(), *simple]


def _check_fields(text: str) -> None:
    """Rechaza los campos cuyo tipo (primera palabra del código) es peligroso o no se puede saber."""
    for code in _field_codes(text):
        if code is None:
            raise TemplateContentError("La plantilla contiene un campo cuyo tipo se calcula con otro campo.")
        words = code.split()
        if words and (field := words[0].upper()) in _DANGEROUS_FIELDS:
            raise TemplateContentError(f"La plantilla contiene un campo {field}, que no se admite.")


def _check_part(name: str, data: bytes) -> None:
    """DTD, relaciones externas que no son hipervínculos y campos peligrosos de una parte XML."""
    text = data.decode("utf-8", "replace")
    if "<!DOCTYPE" in text.upper():
        raise TemplateContentError("La plantilla contiene una DTD (<!DOCTYPE>), que Office no usa.")
    if name.endswith(".rels"):
        for match in _RELATIONSHIP.finditer(text):
            attrs = _attributes(match.group(1))
            if attrs.get("TargetMode", "").lower() == "external" and not attrs.get("Type", "").endswith("/hyperlink"):
                raise TemplateContentError(
                    "La plantilla enlaza ficheros externos (plantilla adjunta, imágenes u objetos enlazados); "
                    "incrústalos o quítalos."
                )
        return
    _check_fields(text)


def _check_package(path: Path, kind: Kind, max_uncompressed: int) -> None:
    if not zipfile.is_zipfile(path):
        raise TemplateContentError("La plantilla no es un documento de Office (no es un ZIP).")
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if sum(info.file_size for info in infos) > max_uncompressed:
            raise TemplateContentError("La plantilla descomprimida es demasiado grande.")
        names = [info.filename for info in infos]
        if _MAIN_PART[kind] not in names:
            raise TemplateContentError(f"El fichero no es un {'Word' if kind == 'word' else 'Excel'}.")
        lowered = [f"/{name.lower()}" for name in names]
        if any(marker in name for name in lowered for marker in _EMBEDDED):
            raise TemplateContentError(
                "La plantilla contiene macros, objetos incrustados (OLE, ActiveX) o vínculos a otros libros."
            )
        for name in names:
            if name.endswith((".xml", ".rels")):
                _check_part(name, archive.read(name))


def _word_info(path: Path) -> TemplateInfo:
    try:
        document = docx.Document(str(path))
    except Exception as exc:  # fichero del usuario: cualquier fallo de lectura es contenido no válido
        raise TemplateContentError("python-docx no puede abrir la plantilla Word.") from exc
    missing = []
    for style in WORD_REQUIRED_STYLES:
        try:
            document.styles[style]
        except KeyError:
            missing.append(style)
    if missing:
        raise TemplateContentError(
            f"A la plantilla le faltan estilos: {', '.join(missing)}. Úsalos una vez en Word (se guardan con el "
            "documento) o parte de la plantilla base.",
            missing_styles=missing,
        )
    styles = sorted({str(s.name) for s in document.styles if s.type == WD_STYLE_TYPE.TABLE and s.name})
    return TemplateInfo(table_styles=styles, table_style="Table Grid" if "Table Grid" in styles else None)


def _excel_info(path: Path) -> TemplateInfo:
    try:
        workbook = load_workbook(path)
    except Exception as exc:  # fichero del usuario
        raise TemplateContentError("openpyxl no puede abrir la plantilla Excel.") from exc
    if not workbook.worksheets:
        raise TemplateContentError("La plantilla Excel no tiene hojas.")
    if not theme_is_safe(workbook.loaded_theme):
        raise TemplateContentError("El tema de la plantilla Excel usa imágenes; elige un tema sin imágenes.")
    return TemplateInfo(table_styles=list(TABLE_STYLES), table_style=DEFAULT_TABLE_STYLE)


def check_template(path: Path, kind: Kind, *, max_uncompressed: int) -> TemplateInfo:
    """Comprueba la plantilla y devuelve sus estilos de tabla.

    Raises:
        TemplateContentError: No es del tipo declarado, es demasiado grande descomprimida, no es segura o (Word) le
            faltan estilos (``missing_styles`` en el *problem details*).
    """
    _check_package(path, kind, max_uncompressed)
    return _word_info(path) if kind == "word" else _excel_info(path)


def base_template(kind: Kind) -> bytes:
    """Plantilla base: Word con todos los estilos que usa util_bbdd y en horizontal; Excel con encabezado y pie."""
    buffer = io.BytesIO()
    if kind == "word":
        document = docx.Document()
        section = document.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        section.header.paragraphs[0].text = "Documentación de base de datos"
        document.save(buffer)
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Plantilla"
        sheet.oddHeader.center.text = "&A"
        sheet.oddFooter.center.text = "Página &P de &N"
        sheet.page_setup.orientation = "landscape"
        workbook.save(buffer)
    return buffer.getvalue()
