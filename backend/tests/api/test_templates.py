"""Plantillas de documento (ADR-011): catálogo global (admins), plantillas de proyecto (editores) y su uso en la
documentación con util_bbdd real sobre SQLite, incluido el formato por defecto con aviso cuando no se pueden aplicar."""

from __future__ import annotations

import hashlib
import io
import json
import zipfile
from collections.abc import Callable
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import Settings
from app.db.models import AuditEntry, DocumentTemplate, User
from app.domain.identity.audit import CLI_CONTEXT
from app.domain.templates import service
from app.integrations.util_bbdd.excel_template import FIXED_COLORS
from app.jobs.worker import Worker
from app.workspace.layout import run_dir
from tests.api.test_documentation import launch, project  # noqa: F401  (fixture)
from tests.api.test_projects import login_as, new_project
from tests.api.test_runs import wait_for, worker  # noqa: F401  (fixture)
from tests.unit.test_templates_units import (
    EXTERNAL_TEMPLATE,
    MARKER,
    add_relationship,
    custom_theme,
    excel_template,
    patch_zip,
    word_template,
)

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def upload(client: TestClient, url: str, filename: str, content: bytes, name: str | None = None) -> Any:
    data = {"name": name} if name is not None else None
    return client.post(url, files={"file": (filename, content, "application/octet-stream")}, data=data)


def project_url(project_id: object) -> str:
    return f"/api/v1/projects/{project_id}/document-templates"


# --- gestión -----------------------------------------------------------------------------------------------------


def test_project_templates_crud_and_roles(
    user_client: TestClient, make_user: Callable[..., User], db: Session, settings: Settings
) -> None:
    project_id = new_project(user_client)["id"]
    url = project_url(project_id)
    created = upload(user_client, url, "../Corporativa.docx", word_template(), name="Corporativa")

    assert created.status_code == 201, created.text
    body = created.json()
    assert (body["scope"], body["kind"], body["name"], body["filename"]) == (
        "project", "word", "Corporativa", "Corporativa.docx"
    )  # fmt: skip
    assert body["table_style"] == "Table Grid" and "path" not in body
    stored = db.get(DocumentTemplate, body["id"])
    assert stored is not None and stored.path.endswith(f"/projects/{project_id}/templates/{body['id']}.docx")
    file = settings.workspace_dir / stored.path
    assert hashlib.sha256(file.read_bytes()).hexdigest() == body["sha256"]

    assert upload(user_client, url, "otra.docx", word_template(), name="corporativa").status_code == 409
    renamed = user_client.patch(f"{url}/{body['id']}", json={"name": "Marca", "table_style": "Light Grid Accent 1"})
    assert renamed.status_code == 200 and renamed.json()["table_style"] == "Light Grid Accent 1"
    assert user_client.patch(f"{url}/{body['id']}", json={"table_style": "Inventado"}).status_code == 422
    assert user_client.patch(f"{url}/{body['id']}", json={"color": "rojo"}).status_code == 422

    viewer = login_as(user_client, make_user, "vera")
    stranger = login_as(user_client, make_user, "sara")
    user_client.post(f"/api/v1/projects/{project_id}/members", json={"username": "vera"})
    assert [t["name"] for t in viewer.get(url).json()] == ["Marca"]
    download = viewer.get(f"{url}/{body['id']}/file")
    assert download.status_code == 200 and download.headers["content-type"] == DOCX
    assert 'attachment; filename="Corporativa.docx"' in download.headers["content-disposition"]
    assert download.headers["x-content-type-options"] == "nosniff"
    assert upload(viewer, url, "v.docx", word_template()).status_code == 403
    assert viewer.delete(f"{url}/{body['id']}").status_code == 403
    assert stranger.get(url).status_code == 404
    assert stranger.get(f"{url}/{body['id']}/file").status_code == 404
    assert TestClient(user_client.app, base_url="https://testserver").get(url).status_code == 401

    assert user_client.delete(f"{url}/{body['id']}").status_code == 204
    assert not file.exists() and user_client.get(url).json() == []
    actions = {e.action for e in db.scalars(select(AuditEntry)).all()}
    assert {"template.create", "template.update", "template.delete"} <= actions


def test_global_templates_are_managed_by_admins(user_client: TestClient, make_user: Callable[..., User]) -> None:
    admin = login_as(user_client, make_user, "root", role="admin")
    assert upload(user_client, "/api/v1/admin/document-templates", "a.xlsx", excel_template()).status_code == 403

    created = upload(admin, "/api/v1/admin/document-templates", "Marca.xltx", excel_template())
    assert created.status_code == 201, created.text
    body = created.json()
    assert (body["scope"], body["kind"], body["name"], body["table_style"]) == (
        "global", "excel", "Marca", "TableStyleMedium2"
    )  # fmt: skip
    assert [t["id"] for t in user_client.get("/api/v1/document-templates?kind=excel").json()] == [body["id"]]
    assert user_client.get("/api/v1/document-templates?kind=word").json() == []
    download = user_client.get(f"/api/v1/document-templates/{body['id']}/file")
    assert download.status_code == 200 and download.headers["content-type"] == XLSX
    assert 'filename="Marca.xlsx"' in download.headers["content-disposition"]

    project_id = new_project(user_client)["id"]
    url = project_url(project_id)
    own = upload(user_client, url, "propia.docx", word_template()).json()
    listed = user_client.get(url).json()
    assert [(t["scope"], t["name"]) for t in listed] == [("project", "propia"), ("global", "Marca")]
    assert user_client.get(f"{url}/{body['id']}/file").status_code == 200
    assert user_client.patch(f"{url}/{body['id']}", json={"name": "Mía"}).status_code == 404
    assert user_client.delete(f"{url}/{body['id']}").status_code == 404
    assert user_client.get(f"/api/v1/document-templates/{own['id']}/file").status_code == 404

    patched = admin.patch(f"/api/v1/admin/document-templates/{body['id']}", json={"table_style": "TableStyleDark1"})
    assert patched.status_code == 200 and patched.json()["table_style"] == "TableStyleDark1"
    assert admin.delete(f"/api/v1/admin/document-templates/{body['id']}").status_code == 204


@pytest.mark.parametrize(
    ("filename", "content", "status", "detail"),
    [
        ("plantilla.pdf", b"%PDF", 415, None),
        ("plantilla.docm", b"x", 415, "macros"),
        ("plantilla.docx", b"no es un zip", 422, "no es un ZIP"),
        ("plantilla.docx", word_template(drop_styles=("Heading 2",)), 422, "Heading 2"),
        (
            "plantilla.docx",
            patch_zip(word_template(), {"word/_rels/document.xml.rels": add_relationship(EXTERNAL_TEMPLATE)}),
            422,
            "ficheros externos",
        ),
        ("plantilla.xlsx", word_template(), 422, "no es un Excel"),
    ],
    ids=["pdf", "docm", "no-zip", "sin-estilos", "externa", "word-como-excel"],
)
def test_uploads_are_checked(
    user_client: TestClient, filename: str, content: bytes, status: int, detail: str | None
) -> None:
    project_id = new_project(user_client)["id"]

    response = upload(user_client, project_url(project_id), filename, content)

    assert response.status_code == status, response.text
    if detail:
        assert detail in response.json()["detail"]
    if "Heading 2" in (detail or ""):
        assert response.json()["missing_styles"] == ["Heading 2"]
    assert user_client.get(project_url(project_id)).json() == []


def _stored(settings: Settings) -> list[str]:
    """Plantillas en disco (sin temporales), de cualquier ámbito."""
    return sorted(p.name for p in settings.workspace_dir.rglob("*") if p.parent.name == "templates" and p.is_file())


def test_concurrent_duplicate_names_are_stopped_by_the_database(
    user_client: TestClient, make_user: Callable[..., User], settings: Settings, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Dos subidas a la vez pasan la comprobación previa: el índice único frena la segunda (también en las globales,
    con ``project_id`` nulo) y su fichero no se queda en disco."""
    monkeypatch.setattr(service, "_check_name", lambda *args, **kwargs: None)  # como si llegaran a la vez
    url = project_url(new_project(user_client)["id"])
    first = upload(user_client, url, "a.docx", word_template(), name="Marca")
    assert first.status_code == 201

    clash = upload(user_client, url, "b.docx", word_template(), name="MARCA")
    assert clash.status_code == 409 and "llamada «MARCA»" in clash.json()["detail"]
    assert upload(user_client, url, "c.xlsx", excel_template(), name="Marca").status_code == 201  # otro tipo
    admin = login_as(user_client, make_user, "root", role="admin")
    assert (
        upload(admin, "/api/v1/admin/document-templates", "g.docx", word_template(), name="Global").status_code == 201
    )
    assert (
        upload(admin, "/api/v1/admin/document-templates", "h.docx", word_template(), name="global").status_code == 409
    )
    assert len(_stored(settings)) == 3

    other = upload(user_client, url, "d.docx", word_template(), name="Otra").json()
    renamed = user_client.patch(f"{url}/{other['id']}", json={"name": "marca"})
    assert renamed.status_code == 409
    assert [t["name"] for t in user_client.get(url).json() if t["kind"] == "word" and t["scope"] == "project"] == [
        "Marca", "Otra"
    ]  # fmt: skip


def test_limit_is_checked_again_after_inserting(
    user_client: TestClient, settings: Settings, db: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Si la comprobación previa del máximo no ve la otra subida, el recuento tras el INSERT la rechaza."""
    original = service._check_limit

    def only_after_insert(*args: Any, pending: int = 1) -> None:
        if pending == 0:
            original(*args, pending=pending)

    monkeypatch.setattr(service, "_check_limit", only_after_insert)
    settings.templates_per_scope = 1
    url = project_url(new_project(user_client)["id"])
    assert upload(user_client, url, "a.docx", word_template()).status_code == 201

    response = upload(user_client, url, "b.docx", word_template())

    assert response.status_code == 409 and "máximo de 1" in response.json()["detail"]
    assert len(db.scalars(select(DocumentTemplate)).all()) == 1
    assert len(_stored(settings)) == 1


def test_file_is_removed_when_the_row_cannot_be_saved(
    db: Session, user: User, settings: Settings, monkeypatch: pytest.MonkeyPatch
) -> None:
    def broken(*args: Any, **kwargs: Any) -> None:
        raise RuntimeError("la BD falló al guardar")

    monkeypatch.setattr(service.audit, "record", broken)

    with pytest.raises(RuntimeError):
        service.save_template(
            db, user, None, name="x", filename="x.docx", stream=io.BytesIO(word_template()), settings=settings,
            ctx=CLI_CONTEXT,
        )  # fmt: skip

    assert _stored(settings) == []
    assert not list(settings.workspace_dir.rglob(".tmp-*"))
    assert db.scalars(select(DocumentTemplate)).all() == []


def test_size_and_count_limits(user_client: TestClient, settings: Settings) -> None:
    project_id = new_project(user_client)["id"]
    url = project_url(project_id)
    settings.template_max_mb = 1
    assert upload(user_client, url, "g.docx", b"0" * (1024 * 1024 + 1)).status_code == 413
    settings.templates_per_scope = 1
    assert upload(user_client, url, "a.docx", word_template()).status_code == 201
    response = upload(user_client, url, "b.docx", word_template())
    assert response.status_code == 409 and "máximo de 1" in response.json()["detail"]
    assert not list((settings.workspace_dir).rglob(".tmp-*"))


@pytest.mark.parametrize("kind", ["word", "excel"])
def test_base_templates_can_be_downloaded_and_uploaded(user_client: TestClient, kind: str) -> None:
    response = user_client.get(f"/api/v1/document-templates/base/{kind}")

    assert response.status_code == 200
    extension = "docx" if kind == "word" else "xlsx"
    assert f'filename="plantilla_base.{extension}"' in response.headers["content-disposition"]
    project_id = new_project(user_client)["id"]
    assert upload(user_client, project_url(project_id), f"base.{extension}", response.content).status_code == 201
    assert user_client.get("/api/v1/document-templates/base/pdf").status_code == 422


def test_catalog_exposes_template_params(user_client: TestClient) -> None:
    operations = {o["name"]: o for o in user_client.get("/api/v1/operations").json()}
    doc = {p["name"]: p for p in operations["documentation"]["params"]}

    assert (doc["word_template"]["type"], doc["word_template"]["template_kind"]) == ("template", "word")
    assert (doc["excel_template"]["group"], doc["excel_template"]["template_kind"]) == ("templates", "excel")
    assert doc["word_template"]["default"] is None and doc["tables"]["template_kind"] is None


# --- uso en la documentación -------------------------------------------------------------------------------------


def _templates(client: TestClient, project_id: object) -> tuple[str, str]:
    url = project_url(project_id)
    word = upload(client, url, "marca.docx", word_template()).json()["id"]
    excel = upload(client, url, "marca.xlsx", excel_template()).json()["id"]
    return word, excel


def _outputs(directory: Any, kind: str) -> list[Any]:
    return sorted((directory / "outputs" / kind).glob(f"*.{'docx' if kind == 'word' else 'xlsx'}"))


def _fixed_fills(book: Any) -> set[str]:
    workbook = load_workbook(book)
    found = set()
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                color = cell.fill.fgColor
                if cell.fill.fill_type == "solid" and color.type == "rgb" and color.rgb[-6:].upper() in FIXED_COLORS:
                    found.add(color.rgb[-6:].upper())
    return found


def test_documentation_applies_word_and_excel_templates(
    user_client: TestClient,
    project: dict[str, object],  # noqa: F811
    worker: Worker,  # noqa: F811
    db: Session,
    settings: Settings,
) -> None:
    word, excel = _templates(user_client, project["id"])
    user_client.patch(f"{project_url(project['id'])}/{excel}", json={"table_style": "TableStyleLight9"})

    run_id = launch(
        user_client, project["id"], "documentation", connection="tienda", graphs=False, html=False,
        word_template=word, excel_template=excel,
    )  # fmt: skip
    run = wait_for(worker, db, run_id, timeout=180)

    assert run.status == "succeeded", (run.error_code, run.error_message)
    directory = run_dir(settings.workspace_dir, str(run.workspace_path))
    with zipfile.ZipFile(_outputs(directory, "word")[0]) as document:
        xml = document.read("word/document.xml").decode("utf-8")
    assert MARKER in xml and "pedidos" in xml
    book = _outputs(directory, "excel")[0]
    workbook = load_workbook(book)
    assert workbook.loaded_theme == custom_theme()
    assert {t.tableStyleInfo.name for s in workbook.worksheets for t in s.tables.values()} == {"TableStyleLight9"}
    assert _fixed_fills(book) == set()  # util_bbdd 4.0.0 no usa otros colores fijos (G40)
    assert (directory / "inputs" / "plantilla_word.docx").is_file()
    manifest = json.loads((directory / "run.json").read_text(encoding="utf-8"))
    assert [(t["kind"], t["applied"], t["scope"]) for t in manifest["templates"]] == [
        ("word", True, "project"), ("excel", True, "project")
    ]  # fmt: skip
    assert manifest["warnings"] == []
    detail = user_client.get(f"/api/v1/runs/{run_id}").json()
    assert detail["warnings"] == [] and not any(a["path"].startswith("inputs/") for a in detail["artifacts"])


def _replace_file(db: Session, settings: Settings, template_id: str, content: bytes) -> None:
    """Simula una plantilla que pasó la comprobación pero falla al usarla (o la de una versión anterior)."""
    template = db.get(DocumentTemplate, template_id)
    assert template is not None
    (settings.workspace_dir / template.path).write_bytes(content)
    template.sha256 = hashlib.sha256(content).hexdigest()
    db.commit()


def test_templates_that_cannot_be_applied_fall_back_with_a_warning(
    user_client: TestClient,
    project: dict[str, object],  # noqa: F811
    worker: Worker,  # noqa: F811
    db: Session,
    settings: Settings,
) -> None:
    word, excel = _templates(user_client, project["id"])
    _replace_file(db, settings, word, word_template(drop_styles=("Heading 1",)))
    _replace_file(db, settings, excel, b"no es un libro")

    run_id = launch(
        user_client, project["id"], "documentation", connection="tienda", graphs=False,
        word_template=word, excel_template=excel,
    )  # fmt: skip
    run = wait_for(worker, db, run_id, timeout=180)

    assert run.status == "succeeded", (run.error_code, run.error_message)
    assert run.progress_message and run.progress_message.endswith("(con avisos)")
    directory = run_dir(settings.workspace_dir, str(run.workspace_path))
    with zipfile.ZipFile(_outputs(directory, "word")[0]) as document:
        xml = document.read("word/document.xml").decode("utf-8")
    assert MARKER not in xml and "pedidos" in xml
    assert _fixed_fills(_outputs(directory, "excel")[0])  # el libro de util_bbdd, sin retocar
    assert (directory / "outputs" / "html" / "documentacion" / "index.html").is_file()
    detail = user_client.get(f"/api/v1/runs/{run_id}").json()
    assert detail["warnings"] == [
        "La plantilla Word no se pudo aplicar: el documento Word usa el formato por defecto.",
        "La plantilla Excel no se pudo aplicar: el libro Excel usa el formato por defecto.",
    ]
    manifest = json.loads((directory / "run.json").read_text(encoding="utf-8"))
    assert [t["applied"] for t in manifest["templates"]] == [False, False]
    log = (directory / "logs" / "run.log").read_text(encoding="utf-8")
    assert "Plantilla Word no aplicada" in log and "Plantilla Excel no aplicada" in log


def test_missing_or_changed_template_files_fall_back(
    user_client: TestClient,
    project: dict[str, object],  # noqa: F811
    worker: Worker,  # noqa: F811
    db: Session,
    settings: Settings,
) -> None:
    word, excel = _templates(user_client, project["id"])
    stored = {t.id: t for t in db.scalars(select(DocumentTemplate)).all()}
    (settings.workspace_dir / stored[word].path).unlink()
    (settings.workspace_dir / stored[excel].path).write_bytes(excel_template(theme=custom_theme() + b" "))

    run_id = launch(
        user_client, project["id"], "documentation", connection="tienda", graphs=False, html=False,
        word_template=word, excel_template=excel,
    )  # fmt: skip
    run = wait_for(worker, db, run_id, timeout=180)

    assert run.status == "succeeded", (run.error_code, run.error_message)
    warnings = user_client.get(f"/api/v1/runs/{run_id}").json()["warnings"]
    assert warnings == [
        "La plantilla Word ya no está disponible: el documento Word usa el formato por defecto.",
        "La plantilla Excel ya no está disponible: el libro Excel usa el formato por defecto.",
    ]
    directory = run_dir(settings.workspace_dir, str(run.workspace_path))
    assert not (directory / "inputs" / "plantilla_word.docx").exists()


def test_runs_only_accept_visible_templates_of_the_right_kind(
    user_client: TestClient,
    project: dict[str, object],  # noqa: F811
) -> None:
    word, _ = _templates(user_client, project["id"])
    other = new_project(user_client, name="Otro")["id"]
    foreign = upload(user_client, project_url(other), "ajena.docx", word_template()).json()["id"]
    url = f"/api/v1/projects/{project['id']}/runs"

    def post(**params: object) -> Any:
        return user_client.post(url, json={"operation": "documentation", "params": {"connection": "tienda", **params}})

    response = post(word_template=foreign)
    assert response.status_code == 422 and "no existe en el proyecto" in response.json()["detail"]
    assert response.json()["errors"][0]["loc"] == ["body", "params", "word_template"]
    assert post(excel_template=word).status_code == 422
    assert post(word_template="no-es-un-ulid").status_code == 422
    accepted = post(word=False, word_template=foreign)  # formato sin generar: la plantilla se descarta
    assert accepted.status_code == 202 and accepted.json()["params"]["word_template"] is None


def test_templates_in_use_cannot_be_deleted(
    user_client: TestClient,
    project: dict[str, object],  # noqa: F811
) -> None:
    word, excel = _templates(user_client, project["id"])
    launch(user_client, project["id"], "documentation", connection="tienda", word_template=word)  # queda en cola

    assert user_client.delete(f"{project_url(project['id'])}/{word}").status_code == 409
    assert user_client.delete(f"{project_url(project['id'])}/{excel}").status_code == 204
