"""Plantillas de documento (ADR-011): comprobación de Word/Excel, conversión de .dotx/.xltx, plantillas base, retoque
del Excel y detección de fallos de la plantilla Word. Los constructores de plantillas los usan también los tests API."""

from __future__ import annotations

import io
import subprocess
import zipfile
from collections.abc import Callable
from pathlib import Path

import docx
import pytest
from docx.opc.exceptions import PackageNotFoundError
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.writer.theme import theme_xml

from app.domain.templates.errors import TemplateContentError, TemplateTypeError
from app.domain.templates.files import base_template, check_template, display_name, kind_for, untemplate
from app.integrations.util_bbdd.documentation import (
    DocumentationOptions,
    build_imprimir,
    word_template_failed,
)
from app.integrations.util_bbdd.excel_template import FIXED_COLORS, ExcelTemplateError, apply_excel_template

MB = 1024 * 1024
MARKER = "PORTADA-UBW"
EXTERNAL_TEMPLATE = (
    '<Relationship Id="rIdX" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
    'attachedTemplate" Target="https://malo.example/plantilla.dotm" TargetMode="External"/>'
)
EXTERNAL_LINK = (
    '<Relationship Id="rIdY" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
    'Target="https://empresa.example" TargetMode="External"/>'
)


# --- constructores -----------------------------------------------------------------------------------------------


def word_template(*, marker: str = MARKER, drop_styles: tuple[str, ...] = ()) -> bytes:
    """Word válido con un párrafo marcador (``drop_styles``: estilos que se quitan para provocar G41)."""
    document = docx.Document()
    document.add_paragraph(marker)
    for name in drop_styles:
        element = document.styles[name].element
        element.getparent().remove(element)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def patch_zip(data: bytes, parts: dict[str, bytes | Callable[[bytes], bytes]]) -> bytes:
    """Copia del ZIP con partes sustituidas (``bytes``), transformadas (función) o añadidas."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as out:
        names = source.namelist()
        for name in names:
            change = parts.get(name)
            content = source.read(name)
            out.writestr(name, change(content) if callable(change) else change if change is not None else content)
        for name, change in parts.items():
            if name not in names and isinstance(change, bytes):
                out.writestr(name, change)
    return buffer.getvalue()


def add_relationship(relationship: str) -> Callable[[bytes], bytes]:
    return lambda xml: xml.replace(b"</Relationships>", relationship.encode() + b"</Relationships>")


def with_field(code: str) -> Callable[[bytes], bytes]:
    run = f"<w:r><w:instrText>{code}</w:instrText></w:r>".encode()
    return lambda xml: xml.replace(b"</w:p>", run + b"</w:p>", 1)


def _char(kind: str) -> str:
    return f'<w:r><w:fldChar w:fldCharType="{kind}"/></w:r>'


def _instr(text: str) -> str:
    return f'<w:r><w:instrText xml:space="preserve">{text}</w:instrText></w:r>'


def complex_field(*pieces: str) -> Callable[[bytes], bytes]:
    """Campo complejo como lo guarda Word (inicio, código en trozos, separador, resultado y fin).

    Un trozo ``"{"``/``"}"`` abre o cierra un campo anidado dentro del código.
    """
    body = "".join(_char("begin") if p == "{" else _char("end") if p == "}" else _instr(p) for p in pieces)
    run = (_char("begin") + body + _char("separate") + "<w:r><w:t>resultado</w:t></w:r>" + _char("end")).encode()
    return lambda xml: xml.replace(b"</w:p>", run + b"</w:p>", 1)


def custom_theme() -> bytes:
    """Tema de Office con ``accent1`` rojo y la fuente del cuerpo en Arial."""
    xml = theme_xml.replace('<a:srgbClr val="4F81BD"/>', '<a:srgbClr val="C00000"/>', 1)
    minor = xml.index("<a:minorFont>")
    xml = xml[:minor] + xml[minor:].replace('typeface="Calibri"', 'typeface="Arial"', 1)
    return xml.encode("utf-8")


def excel_template(*, theme: bytes | None = None) -> bytes:
    """Excel con tema propio, encabezado con imagen (``&G``), pie y página en horizontal."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.oddHeader.left.text = "Empresa&G"
    sheet.oddFooter.right.text = "Página &P"
    sheet.page_setup.orientation = "landscape"
    workbook.loaded_theme = theme or custom_theme()
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def as_office_template(data: bytes, kind: str) -> bytes:
    """El mismo documento con el tipo de parte de un ``.dotx``/``.xltx``."""
    source, target = {
        "word": (b"wordprocessingml.document.main+xml", b"wordprocessingml.template.main+xml"),
        "excel": (b"spreadsheetml.sheet.main+xml", b"spreadsheetml.template.main+xml"),
    }[kind]
    return patch_zip(data, {"[Content_Types].xml": lambda xml: xml.replace(source, target)})


def _write(tmp_path: Path, name: str, data: bytes) -> Path:
    path = tmp_path / name
    path.write_bytes(data)
    return path


# --- tipo y nombre -----------------------------------------------------------------------------------------------


def test_kind_and_display_name() -> None:
    assert kind_for("Corporativa.DOCX") == ("word", False)
    assert kind_for("a.dotx") == ("word", True)
    assert kind_for("a.xltx") == ("excel", True)
    with pytest.raises(TemplateTypeError, match="macros"):
        kind_for("a.docm")
    with pytest.raises(TemplateTypeError):
        kind_for("a.pdf")
    assert display_name("..\\..\\c:/x/Plan\x07tilla.docx") == "Plantilla.docx"


# --- Word --------------------------------------------------------------------------------------------------------


def test_valid_word_template_lists_table_styles(tmp_path: Path) -> None:
    info = check_template(_write(tmp_path, "t.docx", word_template()), "word", max_uncompressed=50 * MB)

    assert info.table_style == "Table Grid"
    assert "Light Shading Accent 1" in info.table_styles


def test_word_template_without_required_styles_is_rejected(tmp_path: Path) -> None:
    path = _write(tmp_path, "t.docx", word_template(drop_styles=("Heading 1", "No Spacing")))

    with pytest.raises(TemplateContentError) as error:
        check_template(path, "word", max_uncompressed=50 * MB)

    assert error.value.extra["missing_styles"] == ["Heading 1", "No Spacing"]
    assert "plantilla base" in str(error.value)


def test_dotx_is_converted_to_docx(tmp_path: Path) -> None:
    path = _write(tmp_path, "t.dotx", as_office_template(word_template(), "word"))
    with pytest.raises(TemplateContentError, match="python-docx"):
        check_template(path, "word", max_uncompressed=50 * MB)

    untemplate(path, "word")

    assert check_template(path, "word", max_uncompressed=50 * MB).table_style == "Table Grid"


@pytest.mark.parametrize(
    ("parts", "message"),
    [
        ({"word/_rels/document.xml.rels": add_relationship(EXTERNAL_TEMPLATE)}, "ficheros externos"),
        ({"word/document.xml": with_field(" DDEAUTO c:\\\\windows\\\\cmd.exe ")}, "campo DDEAUTO"),
        ({"word/document.xml": with_field(" &#68;DE x ")}, "campo DDE"),
        ({"word/document.xml": with_field(' INCLUDEPICTURE "https://x/y.png" ')}, "INCLUDEPICTURE"),
        ({"word/document.xml": complex_field(" DD", "EAUTO c:\\x ")}, "campo DDEAUTO"),  # código en dos trozos
        ({"word/document.xml": complex_field(" ddeauto x ")}, "campo DDEAUTO"),
        ({"word/document.xml": complex_field(' INCLUDE "c:/x.docx" ')}, "campo INCLUDE"),
        ({"word/document.xml": complex_field("{", " QUOTE 68 ", "}", "DE x ")}, "se calcula con otro campo"),
        (
            {
                "word/document.xml": lambda xml: xml.replace(
                    b"</w:p>", b'<w:fldSimple w:instr=" LINK Excel x "/></w:p>', 1
                )
            },
            "campo LINK",
        ),
        ({"word/embeddings/oleObject1.bin": b"OLE"}, "objetos incrustados"),
        ({"word/vbaProject.bin": b"VBA"}, "macros"),
        ({"word/document.xml": lambda xml: xml.replace(b"?>", b'?><!DOCTYPE x [<!ENTITY e "x">]>', 1)}, "DTD"),
    ],
)
def test_unsafe_word_templates_are_rejected(
    tmp_path: Path, parts: dict[str, bytes | Callable[[bytes], bytes]], message: str
) -> None:
    path = _write(tmp_path, "t.docx", patch_zip(word_template(), parts))

    with pytest.raises(TemplateContentError, match=message):
        check_template(path, "word", max_uncompressed=50 * MB)


def test_external_hyperlinks_and_harmless_fields_are_allowed(tmp_path: Path) -> None:
    parts: dict[str, bytes | Callable[[bytes], bytes]] = {
        "word/_rels/document.xml.rels": add_relationship(EXTERNAL_LINK),
        "word/document.xml": with_field(" PAGE "),
    }
    path = _write(tmp_path, "t.docx", patch_zip(word_template(), parts))

    assert check_template(path, "word", max_uncompressed=50 * MB).table_styles


@pytest.mark.parametrize(
    "pieces",
    [
        (" MERGEFIELD Link ",),
        (" MERGEFIELD import_date ",),
        (
            " MERGE",
            "FIELD Link ",
        ),
        (' HYPERLINK "https://empresa.example" ',),
        (" IF ", "{", " MERGEFIELD Link ", "}", ' = "x" "a" "b" '),
    ],
)
def test_only_the_field_type_counts(tmp_path: Path, pieces: tuple[str, ...]) -> None:
    """Las palabras LINK o IMPORT dentro del código (nombres de campos de combinación) no son el tipo del campo."""
    path = _write(tmp_path, "t.docx", patch_zip(word_template(), {"word/document.xml": complex_field(*pieces)}))

    assert check_template(path, "word", max_uncompressed=50 * MB).table_styles


def test_packages_that_are_not_what_they_say(tmp_path: Path) -> None:
    with pytest.raises(TemplateContentError, match="no es un ZIP"):
        check_template(_write(tmp_path, "a.docx", b"texto"), "word", max_uncompressed=MB)
    with pytest.raises(TemplateContentError, match="no es un Word"):
        check_template(_write(tmp_path, "b.docx", excel_template()), "word", max_uncompressed=50 * MB)
    with pytest.raises(TemplateContentError, match="demasiado grande"):
        check_template(_write(tmp_path, "c.docx", word_template()), "word", max_uncompressed=1024)


# --- Excel -------------------------------------------------------------------------------------------------------


def test_valid_excel_template_and_xltx(tmp_path: Path) -> None:
    info = check_template(_write(tmp_path, "t.xlsx", excel_template()), "excel", max_uncompressed=50 * MB)
    assert info.table_style == "TableStyleMedium2" and "TableStyleLight1" in info.table_styles

    path = _write(tmp_path, "t.xltx", as_office_template(excel_template(), "excel"))
    untemplate(path, "excel")
    with zipfile.ZipFile(path) as archive:
        assert b"spreadsheetml.sheet.main+xml" in archive.read("[Content_Types].xml")
    assert check_template(path, "excel", max_uncompressed=50 * MB).table_styles


def test_unsafe_excel_templates_are_rejected(tmp_path: Path) -> None:
    with_image = custom_theme().replace(b"<a:fillStyleLst>", b'<a:fillStyleLst><a:blipFill><a:blip r:embed="rId1"/>'
                                        b"</a:blipFill>", 1)  # fmt: skip
    with pytest.raises(TemplateContentError, match="tema"):
        check_template(_write(tmp_path, "a.xlsx", excel_template(theme=with_image)), "excel", max_uncompressed=50 * MB)
    linked = patch_zip(excel_template(), {"xl/externalLinks/externalLink1.xml": b"<externalLink/>"})
    with pytest.raises(TemplateContentError, match="vínculos"):
        check_template(_write(tmp_path, "b.xlsx", linked), "excel", max_uncompressed=50 * MB)


@pytest.mark.parametrize("kind", ["word", "excel"])
def test_base_templates_pass_the_checks(tmp_path: Path, kind: str) -> None:
    path = _write(tmp_path, f"base.{'docx' if kind == 'word' else 'xlsx'}", base_template(kind))  # type: ignore[arg-type]

    assert check_template(path, kind, max_uncompressed=50 * MB).table_style  # type: ignore[arg-type]


# --- retoque del Excel -------------------------------------------------------------------------------------------


def _generated_like_util_bbdd(path: Path) -> None:
    """Libro con lo que deja ``ExcelGenerator`` de util_bbdd 4.0.0: rellenos fijos, fuentes sin nombre y tablas."""
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Tablas"
    sheet["A1"].font = Font(size=12, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    sheet.append(["Nombre", "Tipo"])
    sheet.append(["clientes", "TABLE"])
    sheet["A3"].fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    sheet["B3"].fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    sheet.merge_cells("C1:D1")
    table = Table(displayName="T1", ref="A2:B3")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    workbook.save(path)


def test_excel_template_retouches_the_generated_workbook(tmp_path: Path) -> None:
    book = tmp_path / "documentacion.xlsx"
    _generated_like_util_bbdd(book)
    template = _write(tmp_path, "t.xlsx", excel_template())

    apply_excel_template(book, template, "TableStyleLight9")

    workbook = load_workbook(book)
    sheet = workbook.active
    assert workbook.loaded_theme == custom_theme()
    assert (sheet["A1"].fill.fgColor.theme, sheet["A1"].fill.fgColor.tint) == (4, -0.25)
    assert (sheet["A3"].fill.fgColor.theme, sheet["A3"].fill.fgColor.tint) == (4, 0.8)
    assert sheet["B3"].fill.fgColor.rgb.endswith("FFC000")  # color que no es de la paleta fija: se queda
    assert (sheet["A1"].font.name, sheet["A1"].font.scheme, sheet["A1"].font.color.rgb) == (
        "Arial",
        "minor",
        "00FFFFFF",
    )
    assert sheet.tables["T1"].tableStyleInfo.name == "TableStyleLight9"
    assert sheet.oddHeader.left.text == "Empresa" and sheet.oddFooter.right.text == "Página &P"
    assert sheet.page_setup.orientation == "landscape"
    assert not list(tmp_path.glob(".tmp-*"))


def test_unknown_table_style_falls_back_and_broken_template_leaves_the_book(tmp_path: Path) -> None:
    book = tmp_path / "documentacion.xlsx"
    _generated_like_util_bbdd(book)
    apply_excel_template(book, _write(tmp_path, "t.xlsx", excel_template()), "Table Grid")
    assert load_workbook(book).active.tables["T1"].tableStyleInfo.name == "TableStyleMedium2"

    before = book.read_bytes()
    with pytest.raises(Exception):  # noqa: B017 - cualquier fallo de openpyxl vale: lo que importa es el libro
        apply_excel_template(book, _write(tmp_path, "roto.xlsx", b"no es un libro"), None)
    assert book.read_bytes() == before
    with_image = custom_theme().replace(b"<a:fillStyleLst>", b'<a:fillStyleLst><a:blip r:embed="rId1"/>', 1)
    with pytest.raises(ExcelTemplateError):
        apply_excel_template(book, _write(tmp_path, "img.xlsx", excel_template(theme=with_image)), None)
    assert book.read_bytes() == before


def test_fixed_colors_are_the_office_palette() -> None:
    assert set(FIXED_COLORS) == {"4472C4", "4F81BD", "366092", "5B9BD5", "2F75B5", "D9E1F2"}


# --- plantilla Word en util_bbdd ---------------------------------------------------------------------------------


def test_build_imprimir_only_adds_the_template_when_there_is_one(tmp_path: Path) -> None:
    options = DocumentationOptions(formats=("word",))

    assert "plantilla_word" not in build_imprimir(options) and "table_style" not in build_imprimir(options)
    imprimir = build_imprimir(options, word_template=tmp_path / "p.docx", table_style="Light Grid")
    assert (imprimir["plantilla_word"], imprimir["table_style"]) == (str(tmp_path / "p.docx"), "Light Grid")


def _wrapped(inner: BaseException) -> RuntimeError:
    try:
        try:
            raise inner
        except BaseException as exc:
            raise RuntimeError("Error generando la documentación word") from exc
    except RuntimeError as outer:
        return outer


def test_word_template_failures_are_recognized() -> None:
    assert word_template_failed(_wrapped(KeyError("no style with name 'Heading 1'")))
    assert word_template_failed(_wrapped(zipfile.BadZipFile("File is not a zip file")))
    assert word_template_failed(_wrapped(PackageNotFoundError("Package not found at 'p.docx'")))
    assert not word_template_failed(_wrapped(KeyError("tabla")))
    assert not word_template_failed(_wrapped(subprocess.CalledProcessError(1, ["dot"])))
    assert not word_template_failed(RuntimeError("conexión rechazada"))
